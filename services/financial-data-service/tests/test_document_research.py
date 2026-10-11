import asyncio
import base64
import io
import zipfile
import pytest
from pptx import Presentation
from app.documents import AcquisitionRequest, acquire, cvm_company_rows, get_bytes
from app.research_deck import ResearchDeck,build_research_deck

def test_cvm_scope_headers_versions_and_safe_archive():
    raw=io.BytesIO()
    with zipfile.ZipFile(raw,'w') as z:
        z.writestr('report.csv','CNPJ_CIA;VERSAO;MOEDA;VALOR\n12.345.678/0001-90;1;BRL;10\n12.345.678/0001-90;2;BRL;11\n99.999.999/0001-99;1;BRL;999\n')
    doc=cvm_company_rows(raw.getvalue(),'12345678000190',2025,'ITR','https://dados.cvm.gov.br/fixture.zip')
    text=base64.b64decode(doc['base64']).decode()
    assert 'VERSAO;MOEDA;VALOR' in text and ';2;BRL;11' in text and '999' not in text

def test_ssrf_rejected_before_transport():
    for url in ['http://127.0.0.1','https://www.sec.gov.evil.test/x','https://user@data.sec.gov/x','https://www.sec.gov:443/x']:
        with pytest.raises(ValueError):asyncio.run(get_bytes(None,url))

def test_sec_acquisition_exact_issuer_and_no_unbounded_tools(monkeypatch):
    import json
    monkeypatch.setenv('SEC_CONTACT_EMAIL','test@example.test')
    urls=[]
    async def fetch(client,url,limit=8_000_000):
        urls.append(url)
        if 'submissions' in url:return json.dumps({'filings':{'recent':{'form':['10-K'],'accessionNumber':['0000000001-25-000001'],'primaryDocument':['report.htm'],'reportDate':['2025-12-31'],'filingDate':['2026-01-01']}}}).encode()
        return b'<html>Public source text</html>'
    result=asyncio.run(acquire(AcquisitionRequest(jurisdiction='us',registryId='0000000001',startYear=2025,endYear=2025),fetch))
    assert len(result['documents'])==1
    assert urls[-1]=='https://www.sec.gov/Archives/edgar/data/1/000000000125000001/report.htm'
    assert result['documents'][0]['provenance']=='regulator_download'

def test_research_deck_keeps_exact_citations_and_rejects_invented_quotes():
    payload=dict(company='Synthetic',question='What evidence challenges the thesis?',status='needs_human_review',language='en',
      findings=[dict(text='Supplier dependence is a hypothesis.',falsifier='Diversified sourcing could refute it.',citations=[dict(evidenceId='test',quote='Supplier concentration')])],
      evidence=[dict(id='test',text='Supplier concentration was disclosed.',sourceUrl='https://example.test/source',sourceKind='user_upload')],gaps=['Narrative completeness is unresolved.'])
    raw=build_research_deck(ResearchDeck(**payload));deck=Presentation(io.BytesIO(raw))
    text=' '.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert 'Supplier concentration' in text and 'https://example.test/source' in text and 'Human review required' in text
    payload['findings'][0]['citations'][0]['quote']='Fabricated quote'
    with pytest.raises(ValueError):ResearchDeck(**payload)
