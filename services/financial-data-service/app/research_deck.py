"""Original bounded research export: typed plan, deterministic renderer, package QA.
Uses the existing python-pptx stack; does not copy unlicensed ppt-agent source.
"""
from io import BytesIO
from pydantic import BaseModel, Field, model_validator
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


class Citation(BaseModel):
    evidenceId: str = Field(max_length=200)
    quote: str = Field(min_length=10, max_length=1500)


class Finding(BaseModel):
    text: str = Field(max_length=700)
    falsifier: str = Field(max_length=700)
    citations: list[Citation] = Field(min_length=1, max_length=4)


class Evidence(BaseModel):
    id: str = Field(max_length=200)
    text: str = Field(max_length=10000)
    sourceUrl: str = Field(max_length=2048)
    sourceKind: str = Field(max_length=80)
    locator: str | None = Field(default=None,max_length=600)


class ResearchDeck(BaseModel):
    company: str = Field(min_length=1,max_length=160)
    question: str = Field(min_length=20,max_length=1600)
    status: str = Field(max_length=60)
    language: str = Field(pattern=r"^(en|pt)$")
    findings: list[Finding] = Field(max_length=6)
    evidence: list[Evidence] = Field(max_length=90)
    gaps: list[str] = Field(max_length=50)

    @model_validator(mode="after")
    def grounded(self):
        records={e.id:e for e in self.evidence}
        if len(records)!=len(self.evidence): raise ValueError("duplicate_evidence")
        for finding in self.findings:
            for c in finding.citations:
                if c.evidenceId not in records or c.quote not in records[c.evidenceId].text:
                    raise ValueError("citation_invalid")
        if any(len(g)>1500 for g in self.gaps): raise ValueError("gap_too_long")
        return self


def build_research_deck(req: ResearchDeck):
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    pt=req.language=="pt"
    def add(title,body,notes=""):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(8,16,29)
        for text,y,size,color in [(title,.55,28,(248,250,252)),(body,1.55,20,(210,224,240)),
                                  ("GPI × FilingLens · "+("Revisão humana obrigatória" if pt else "Human review required"),6.95,11,(34,211,238))]:
            box=slide.shapes.add_textbox(Inches(.7),Inches(y),Inches(11.9),Inches(.8 if y!=1.55 else 5.1))
            box.text_frame.word_wrap=True
            box.text_frame.text=text
            for p in box.text_frame.paragraphs:
                p.font.name="Aptos";p.font.size=Pt(size);p.font.color.rgb=RGBColor(*color)
        slide.notes_slide.notes_text_frame.text=notes
    add(req.company,("Pesquisa fundamentada" if pt else "Evidence-based research")+"\n\n"+req.status)
    # Long content is paginated rather than silently discarded.
    def pages(title,text,notes=""):
        for i in range(0,len(text),850):add(title,text[i:i+850],notes)
    pages("Pergunta" if pt else "Research question",req.question)
    records={e.id:e for e in req.evidence}
    for i,f in enumerate(req.findings):
        citations="\n\n".join(f"{c.evidenceId}\n{records[c.evidenceId].sourceKind}\n{records[c.evidenceId].sourceUrl}\n{records[c.evidenceId].locator or ''}\nQUOTE: {c.quote}" for c in f.citations)
        pages(f"{i+1}. "+("Resultado para revisão" if pt else "Finding for review"),f.text+"\n\n"+("Teste de refutação: " if pt else "Falsifier: ")+f.falsifier,citations)
        pages("Evidências" if pt else "Evidence and citations",citations,citations)
    if not req.findings:add("Sem conclusão validada" if pt else "No reviewed finding","Evidence is missing, models are disabled, or review withheld the findings.")
    for i in range(0,len(req.gaps),3):pages("Lacunas e limites" if pt else "Gaps and limitations","\n\n".join(req.gaps[i:i+3]))
    add("Método" if pt else "Method","Bounded manager → scoped retrieval → analyst → reviewer. Exact quotes establish linkage, not semantic truth. Numerical valuation remains in FilingLens. No investment approval is implied.")
    output=BytesIO();prs.save(output);raw=output.getvalue()
    reopened=Presentation(BytesIO(raw))
    if len(reopened.slides)!=len(prs.slides):raise ValueError("deck_validation_failed")
    return raw
